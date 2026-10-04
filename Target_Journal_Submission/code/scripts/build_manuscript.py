"""Build the GENETICS manuscript (Markdown + DOCX) from the text sources and machine-generated numbers.

* manuscript/src/*.md hold the text with {{token}} placeholders.
* results/key_numbers.json (from scripts/key_numbers.py) supplies every number.
* Tables are inserted from tables/*.csv at the end of the main text (GENETICS: tables after the main text,
  editable, not images). Figure legends with "Alt text:" lines follow the tables; figures are embedded for
  review convenience and are also supplied as separate vector PDFs.
Unresolved tokens abort the build.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from kincp.utils import FIG_DIR, PKG_ROOT, RESULTS_DIR, TAB_DIR  # noqa: E402

MS = PKG_ROOT / "manuscript"
TOKEN = re.compile(r"\{\{([A-Za-z0-9_]+)\}\}")


def fill(text: str, nums: dict) -> str:
    missing = sorted({m for m in TOKEN.findall(text) if m not in nums})
    if missing:
        raise KeyError(f"unresolved tokens: {missing}")
    return TOKEN.sub(lambda m: str(nums[m.group(1)]), text)


# ------------------------------------------------------------------ inline markdown -> runs
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(par, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Courier New"
        else:
            r = par.add_run(part)
        if size:
            r.font.size = Pt(size)


def line_numbering(section):
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sp = section._sectPr
    # schema order (CT_SectPr): ... pgMar, paperSrc, pgBorders, lnNumType, pgNumType, cols, ... docGrid
    anchor = sp.find(qn("w:pgNumType"))
    if anchor is None:
        anchor = sp.find(qn("w:cols"))
    if anchor is None:
        anchor = sp.find(qn("w:docGrid"))
    if anchor is not None:
        anchor.addprevious(ln)
    else:
        sp.append(ln)


def page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        r._r.append(el)


def set_cell_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tblPr.append(borders)


def add_table(doc, caption, df, font=7.5):
    p = doc.add_paragraph()
    add_runs(p, caption)
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(t)
    for i, c in enumerate(df.columns):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(c))
        r.bold = True
        r.font.size = Pt(font)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row):
            if isinstance(v, float):
                v = f"{v:.3f}" if abs(v) >= 1e-3 or v == 0 else f"{v:.2e}"
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font)
    doc.add_paragraph()


def md_to_docx(md: str, doc: Document, fig_paths: dict):
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln:
            i += 1
            continue
        if ln.startswith("{{FIGURE:"):
            key = ln[len("{{FIGURE:"):-2]
            doc.add_picture(str(fig_paths[key]), width=Inches(6.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if ln.startswith("{{PAGEBREAK}}"):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue
        if ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            text = ln[level:].strip()
            if level == 1:
                p = doc.add_paragraph()
                r = p.add_run(text)
                r.bold = True
                r.font.size = Pt(16)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                h = doc.add_heading(level=min(level - 1, 3))
                add_runs(h, text)
            i += 1
            continue
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [[c.strip() for c in r.strip("|").split("|")] for r in block if not re.match(r"^\|[\s:|-]+\|$", r)]
            t = doc.add_table(rows=0, cols=len(rows[0]))
            set_cell_borders(t)
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                for ci, v in enumerate(row):
                    cells[ci].text = ""
                    add_runs(cells[ci].paragraphs[0], v, size=8)
                    if ri == 0:
                        for r in cells[ci].paragraphs[0].runs:
                            r.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"^(\d+)\. (.*)", ln)
        if m or ln.startswith("* ") or ln.startswith("- "):
            style = "List Number" if m else "List Bullet"
            text = m.group(2) if m else ln[2:]
            p = doc.add_paragraph(style=style)
            add_runs(p, text)
            i += 1
            continue
        # equation line: ends with (n)
        if re.search(r"\s\((\d+)\)\s*$", ln) and len(ln) < 220:
            p = doc.add_paragraph()
            add_runs(p, ln)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        # paragraph (merge consecutive non-empty lines)
        buf = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||\{\{|\* |- |\d+\. )", lines[i]):
            buf.append(lines[i].rstrip())
            i += 1
        p = doc.add_paragraph()
        add_runs(p, " ".join(buf))


def main():
    nums = json.load(open(RESULTS_DIR / "key_numbers.json"))
    parts = [fill((MS / "src" / f).read_text(), nums) for f in sorted(p.name for p in (MS / "src").glob("*.md"))]
    md = "\n\n".join(parts)
    # tables and figure legends
    tabspec = json.load(open(MS / "src" / "tables.json"))
    figspec = json.load(open(MS / "src" / "figures.json"))
    md_full = md + "\n\n## Tables\n\n"
    for t in tabspec:
        df = pd.read_csv(TAB_DIR / t["file"])
        if t.get("query"):
            df = df.query(t["query"])
        if t.get("columns"):
            df = df[t["columns"]]
        md_full += f"**{t['label']}** {fill(t['caption'], nums)}\n\n" + df.to_markdown(index=False, floatfmt=".3f") + "\n\n"
    md_full += "## Figure legends\n\n"
    for f in figspec:
        md_full += f"**{f['label']}** {fill(f['legend'], nums)}\n\nAlt text: {fill(f['alt'], nums)}\n\n"
    (MS / "Final_Manuscript.md").write_text(md_full)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.paragraph_format.line_spacing = 2.0
    st.paragraph_format.space_after = Pt(0)
    for s in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[s].font.name = "Times New Roman"
        doc.styles[s].font.color.rgb = None
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.styles["Heading 3"].font.size = Pt(11)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    line_numbering(sec)
    page_numbers(sec)
    md_to_docx(md, doc, {})
    # tables at end of main text
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_heading(level=1)
    h.add_run("Tables")
    for t in tabspec:
        df = pd.read_csv(TAB_DIR / t["file"])
        if t.get("query"):
            df = df.query(t["query"])
        if t.get("columns"):
            df = df[t["columns"]]
        add_table(doc, f"**{t['label']}** {fill(t['caption'], nums)}", df, font=t.get("font", 7.5))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_heading(level=1)
    h.add_run("Figures and figure legends")
    for f in figspec:
        doc.add_picture(str(FIG_DIR / f"{f['file']}.png"), width=Inches(6.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        add_runs(p, f"**{f['label']}** {fill(f['legend'], nums)}")
        p = doc.add_paragraph()
        add_runs(p, f"Alt text: {fill(f['alt'], nums)}")
        doc.add_paragraph()
    out = MS / "Final_Manuscript.docx"
    doc.save(out)
    print("wrote", out, "and", MS / "Final_Manuscript.md")
    build_cover_letter(nums)


def build_cover_letter(nums: dict):
    """Cover letter (single-spaced business letter) from manuscript/src/cover_letter.md."""
    text = fill((MS / "src" / "cover_letter.md").read_text(), nums)
    (MS / "cover_letter.md").write_text(text)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    for block in text.split("\n\n"):
        p = None
        for ln in block.split("\n"):
            if not ln.strip():
                continue
            if ln.startswith("* "):
                add_runs(doc.add_paragraph(style="List Bullet"), ln[2:])
                p = None
            else:
                if p is None:
                    p = doc.add_paragraph()
                else:
                    p.add_run().add_break()
                add_runs(p, ln)
    doc.save(MS / "cover_letter.docx")
    print("wrote", MS / "cover_letter.docx")


if __name__ == "__main__":
    main()
